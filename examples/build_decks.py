#!/usr/bin/env python3
# /// script
# requires-python = ">=3.10"
# dependencies = ["python-pptx>=1.0.0"]
# ///
"""Rebuild the two original native examples. Run: uv run examples/build_decks.py."""
import argparse
import hashlib
import json
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
INK, BLUE, PALE, MUTED = "183040", "175BC2", "EEF4FA", "526574"


def rgb(value):
    return RGBColor.from_string(value)


def style_run(run, size, color=INK, bold=False, font="PingFang SC"):
    run.font.size, run.font.name = Pt(size), font
    run.font.bold, run.font.color.rgb = bold, rgb(color)
    props = run._r.get_or_add_rPr()
    props.set("lang", "zh-CN")
    ea = props.find("{http://schemas.openxmlformats.org/drawingml/2006/main}ea")
    if ea is None:
        ea = OxmlElement("a:ea")
        props.append(ea)
    ea.set("typeface", font)


def text(slide, value, x, y, w, h, size=22, color=INK, bold=False, role="body"):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = "role=" + role
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = Inches(0)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    for index, line in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.line_spacing = 1.2
        p.space_after = Pt(10)
        style_run(p.add_run(), size, color, bold)
        p.runs[0].text = line
    return shape


def base(prs, title, number, case):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb("FFFFFF")
    text(slide, title, .6, .45, 12.1, 1.0, 36, bold=True, role="slide-title")
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(.6), Inches(6.9), Inches(12.7), Inches(6.9))
    line.line.color.rgb = rgb("D5E0E8")
    text(slide, case + " · 原创演示案例 · 字体未嵌入", .6, 7.02, 10, .25, 11, MUTED, role="footnote")
    text(slide, str(number), 12, 7.02, .7, .25, 11, MUTED, role="page-number")
    return slide


def table(slide, rows):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(.6), Inches(2.1), Inches(12.1), Inches(3.5))
    shape.name = "role=table"
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = shape.table.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(BLUE if i == 0 else (PALE if i % 2 else "FFFFFF"))
            cell.margin_left = Inches(.18)
            cell.margin_top = Inches(.14)
            cell.text = str(value)
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    style_run(run, 20, "FFFFFF" if i == 0 else INK, i == 0)
    return shape


def build(case):
    folder = ROOT / case
    data = json.loads((folder / "source.json").read_text())
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
    for n, item in enumerate(data["slides"], 1):
        slide = base(prs, item["title"], n, data["name"])
        kind = item["kind"]
        if kind == "cover":
            # Preserve actual title as a native object; spacious cover body.
            text(slide, item["headline"], .6, 2, 11.8, 2.1, 48, BLUE, True, "cover-title")
            text(slide, item["body"], .6, 4.7, 11.7, 1.4, 24)
        elif kind == "table":
            text(slide, item["lead"], .6, 1.55, 12, .5, 24, BLUE, role="key-message")
            table(slide, item["rows"])
            text(slide, item["body"], .6, 5.9, 12, .6, 20, MUTED)
        elif kind == "chart":
            text(slide, item["lead"], .6, 1.55, 12, .5, 24, BLUE, role="key-message")
            chart_data = CategoryChartData()
            chart_data.categories = item["categories"]
            chart_data.add_series(item["series"], item["values"])
            chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(.6), Inches(2.15), Inches(8), Inches(3.95), chart_data).chart
            chart.has_legend = False
            chart.has_title = False
            chart.font.name, chart.font.size = "PingFang SC", Pt(18)
            chart.value_axis.maximum_scale = 1
            chart.value_axis.minimum_scale = 0
            chart.value_axis.tick_labels.number_format = '0%'
            chart.category_axis.tick_labels.font.size = Pt(18)
            chart.value_axis.tick_labels.font.size = Pt(18)
            plot = chart.plots[0]
            plot.has_data_labels = True
            plot.data_labels.number_format = '0%'
            plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
            plot.data_labels.font.size = Pt(18)
            plot.series[0].format.fill.solid()
            plot.series[0].format.fill.fore_color.rgb = rgb(BLUE)
            text(slide, item["body"], 9, 2.4, 3.65, 3.6, 22)
        elif kind == "flow":
            text(slide, item["lead"], .6, 1.55, 12, .5, 24, BLUE, role="key-message")
            labels = item["labels"]
            width = (12.1 - .45 * (len(labels) - 1)) / len(labels)
            for i, label in enumerate(labels):
                x = .6 + i * (width + .45)
                node = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(2.65), Inches(width), Inches(1.1))
                node.name = f"role=node;group=flow;index={i}"
                node.fill.solid()
                node.fill.fore_color.rgb = rgb(PALE)
                node.line.fill.background()
                text(slide, label, x + .1, 2.87, width - .2, .6, 20, BLUE, True, "diagram-node")
                if i:
                    edge = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x - .45), Inches(3.2), Inches(x), Inches(3.2))
                    edge.name = "role=connector;group=flow"
                    edge.line.color.rgb = rgb(BLUE)
                    arrow = OxmlElement("a:tailEnd")
                    arrow.set("type", "triangle")
                    edge.line._get_or_add_ln().append(arrow)
            text(slide, item["body"], .6, 4.35, 12, 1.7, 24)
        else:
            text(slide, item["lead"], .6, 1.65, 12, 1.0, 28, BLUE, role="key-message")
            for i, (title, body) in enumerate(item["sections"]):
                y = 3.0 + i * 1.15
                text(slide, title, .6, y, 3, .6, 22, bold=True, role="module-title")
                text(slide, body, 3.9, y, 8.8, .85, 22)
        slide.notes_slide.notes_text_frame.text = item["notes"]
    destination = folder / "deck.pptx"
    prs.save(destination)
    # Structural evidence refers to the saved file, not the in-memory object.
    actual = Presentation(destination)
    receipt = {
        "file": "deck.pptx", "sha256": hashlib.sha256(destination.read_bytes()).hexdigest(),
        "bytes": destination.stat().st_size, "slides": len(actual.slides),
        "native_tables": sum(s.has_table for p in actual.slides for s in p.shapes),
        "native_charts": sum(s.has_chart for p in actual.slides for s in p.shapes),
        "native_text_shapes": sum(s.has_text_frame for p in actual.slides for s in p.shapes),
        "notes": len(actual.slides), "font": "PingFang SC", "font_embedded": False,
        "builder": "python-pptx", "authoring": "Codex single worked run using local subaru-slides guidance",
        "data_boundary": data["boundary"],
        "verified": ["export and re-import", "native object counts"],
        "not_verified": ["PowerPoint", "Keynote", "Google Slides", "recipient fonts", "independent design review"]
    }
    (folder / "receipt.json").write_text(json.dumps(receipt, ensure_ascii=False, indent=2) + "\n")
    print(f"{case}: {len(actual.slides)} slides, {destination.stat().st_size} bytes")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--case", choices=["management-report", "rag-sharing", "all"], default="all")
    args = parser.parse_args()
    for case in ([args.case] if args.case != "all" else ["management-report", "rag-sharing"]):
        build(case)


if __name__ == "__main__":
    main()
